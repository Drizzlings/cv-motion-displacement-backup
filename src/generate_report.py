import argparse
import csv
from pathlib import Path
from typing import Dict, List

import yaml
from docx import Document
from docx.shared import Inches


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = PROJECT_ROOT / "config.yaml"


def load_config(path: Path) -> dict:
    with path.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def project_path(value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else PROJECT_ROOT / path


def read_results(path: Path) -> List[Dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def markdown_table(rows: List[Dict[str, str]]) -> str:
    if not rows:
        return ""
    headers = list(rows[0].keys())
    table = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        table.append("| " + " | ".join(str(row.get(h, "")) for h in headers) + " |")
    return "\n".join(table)


def build_markdown_report(rows: List[Dict[str, str]], config: dict, output_path: Path) -> None:
    lines = [
        "# 运动物体位移测量实验报告",
        "",
        "## 1. 实验目标",
        "",
        "开发一个基于经典图像处理与几何视觉的系统，从固定相机拍摄的平面运动视频中测量运动物体的真实物理位移。",
        "",
        "## 2. 数据与标定",
        "",
        f"- 输入视频目录：`{config['input_dir']}`",
        f"- 标定点文件：`{config['calibration']['points_file']}`",
        f"- 标定参照：纸上内框，尺寸为 `{config['calibration']['reference_width_mm']}mm x {config['calibration']['reference_height_mm']}mm`。",
        "- 标定方式：用户在视频帧上手动标出内框四个角点，顺序为左上、右上、右下、左下。",
        "",
        "## 3. 算法流程",
        "",
        "1. 用户点击内框四角，程序读取四点并计算单应性矩阵。",
        "2. 使用 `cv2.findHomography` 将图像坐标转换到毫米坐标。",
        "3. 用户在初始帧框选运动目标区域。",
        "4. 在目标框内提取 Shi-Tomasi 角点，并使用 Lucas-Kanade 光流逐帧跟踪。",
        "5. 使用特征点的中位位移更新目标中心。",
        "6. 将目标中心映射到毫米坐标，计算直线位移和累计路径长度。",
        "",
        "## 4. 参数设置",
        "",
        f"- 目标框文件：`{config['optical_flow']['target_rois_file']}`",
        f"- 最大角点数：`{config['optical_flow']['max_corners']}`",
        f"- 角点质量阈值：`{config['optical_flow']['quality_level']}`",
        f"- LK 窗口大小：`{config['optical_flow']['win_size']}`",
        f"- 金字塔层数：`{config['optical_flow']['max_level']}`",
        f"- 最少有效跟踪点：`{config['optical_flow']['min_tracks']}`",
        f"- 最大帧间中心跳变：`{config['optical_flow']['max_center_jump_px']}` 像素",
        f"- 最大物理步长过滤：`{config['quality']['max_world_step_mm']}` mm/帧",
        f"- 平滑窗口：`{config['optical_flow']['smooth_window']}` 帧",
        "",
        "## 5. 实验结果",
        "",
        markdown_table(rows),
        "",
        "## 6. 结果说明",
        "",
        "测量程序为每段视频生成标注视频、标定图、位移曲线和逐帧 CSV 数据。标注视频中的 `Disp` 表示相对起点的直线位移，`Path` 表示累计运动路径长度。",
        "",
        "当前数据集中没有人工标注的真实位移，因此本报告不计算误差百分比。若后续提供真实位移，可用 `abs(测量值 - 真值) / 真值 * 100%` 计算误差。",
        "",
        "## 7. 局限性",
        "",
        "- 标定精度依赖手动点击的四个角点和内框真实尺寸。",
        "- 光流跟踪依赖初始目标框和可跟踪角点质量，透明带或模糊会降低稳定性。",
        "- 若相机发生明显移动或目标离开标定平面，真实位移会产生误差。",
    ]
    output_path.write_text("\n".join(lines), encoding="utf-8")


def build_docx_report(rows: List[Dict[str, str]], config: dict, output_path: Path, output_dir: Path) -> None:
    doc = Document()
    doc.add_heading("运动物体位移测量实验报告", 0)
    doc.add_heading("1. 实验目标", level=1)
    doc.add_paragraph("开发一个基于经典图像处理与几何视觉的系统，从固定相机拍摄的平面运动视频中测量运动物体的真实物理位移。")

    doc.add_heading("2. 数据与标定", level=1)
    doc.add_paragraph(f"输入视频目录：{config['input_dir']}")
    doc.add_paragraph(f"标定点文件：{config['calibration']['points_file']}")
    doc.add_paragraph(
        f"标定参照为纸上内框，真实尺寸为 {config['calibration']['reference_width_mm']}mm x "
        f"{config['calibration']['reference_height_mm']}mm。用户按左上、右上、右下、左下顺序手动点击四个角点，程序使用单应性矩阵完成像素坐标到毫米坐标的转换。"
    )

    doc.add_heading("3. 算法流程", level=1)
    for item in [
        "用户点击内框四角，程序读取四点并计算单应性矩阵。",
        "使用 cv2.findHomography 将图像坐标转换到毫米坐标。",
        "用户在初始帧框选运动目标区域。",
        "在目标框内提取 Shi-Tomasi 角点，并使用 Lucas-Kanade 光流逐帧跟踪。",
        "使用特征点的中位位移更新目标中心。",
        "将目标中心映射到毫米坐标，计算直线位移和累计路径长度。",
    ]:
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("4. 实验结果", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["视频", "帧数", "有效帧", "有效率", "最终位移(mm)", "最大位移(mm)", "状态"]
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = header
    for row in rows:
        cells = table.add_row().cells
        values = [
            row.get("video", ""),
            row.get("frames", ""),
            row.get("valid_frames", ""),
            row.get("valid_ratio", ""),
            row.get("final_displacement_mm", ""),
            row.get("max_displacement_mm", ""),
            row.get("status", ""),
        ]
        for cell, value in zip(cells, values):
            cell.text = value

    doc.add_heading("5. 曲线与标定示例", level=1)
    for row in rows:
        stem = Path(row.get("video", "")).stem
        plot = output_dir / "plots" / f"{stem}_displacement.png"
        calib = output_dir / "calibration" / f"{stem}_calibration.jpg"
        doc.add_paragraph(row.get("video", ""))
        if calib.exists():
            doc.add_picture(str(calib), width=Inches(5.5))
        if plot.exists():
            doc.add_picture(str(plot), width=Inches(5.5))

    doc.add_heading("6. 说明与局限性", level=1)
    doc.add_paragraph("当前数据集中没有人工标注的真实位移，因此本报告不计算误差百分比。若后续提供真实位移，可追加误差统计。")
    doc.add_paragraph("标定精度依赖手动点击的角点和内框真实尺寸；光流跟踪依赖初始目标框和可跟踪角点质量。若相机发生移动或目标离开标定平面，测量结果会产生误差。")
    doc.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Markdown and Word reports from measurement outputs.")
    parser.add_argument("--config", default=str(DEFAULT_CONFIG), help="Path to YAML configuration file.")
    args = parser.parse_args()

    config = load_config(Path(args.config))
    output_dir = project_path(config["output_dir"])
    results_path = output_dir / "results.csv"
    if not results_path.exists():
        raise FileNotFoundError(f"missing results file: {results_path}")
    rows = read_results(results_path)
    build_markdown_report(rows, config, output_dir / "report.md")
    build_docx_report(rows, config, output_dir / "report.docx", output_dir)
    print(f"[INFO] Reports written to {output_dir / 'report.md'} and {output_dir / 'report.docx'}")


if __name__ == "__main__":
    main()
